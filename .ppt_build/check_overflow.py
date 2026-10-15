from math import ceil
import re
import sys

from pptx import Presentation
from pptx.util import Emu


EMU_PER_INCH = 914400
PX_PER_INCH = 96


def east_asian_width(ch):
    return 1.0 if ord(ch) >= 0x2E80 or ch in "，。、：；！？“”‘’（）《》—…·" else 0.55


def estimate_line_height(paragraph, default_pt=18.0, fallback_leading=1.3):
    size_pt = None
    for run in paragraph.runs:
        if run.font.size is not None:
            size_pt = run.font.size.pt
            break
    if size_pt is None:
        size_pt = default_pt
    leading = fallback_leading
    ls = paragraph.line_spacing
    if ls is None:
        leading = fallback_leading
    elif isinstance(ls, float):
        leading = ls
    else:
        try:
            leading = ls.pt / size_pt
        except Exception:
            leading = fallback_leading
    return size_pt * PX_PER_INCH / 72.0 * leading


def estimate_paragraph_width(paragraph, default_pt=18.0):
    size_pt = None
    for run in paragraph.runs:
        if run.font.size is not None:
            size_pt = run.font.size.pt
            break
    if size_pt is None:
        size_pt = default_pt
    char_px = size_pt * PX_PER_INCH / 72.0
    total = 0.0
    for ch in paragraph.text:
        total += east_asian_width(ch) * char_px
    return total


def flag_shapes(path):
    prs = Presentation(path)
    flags = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            if not tf.text.strip():
                continue
            left = tf.margin_left if tf.margin_left is not None else 0
            right = tf.margin_right if tf.margin_right is not None else 0
            top = tf.margin_top if tf.margin_top is not None else 0
            bottom = tf.margin_bottom if tf.margin_bottom is not None else 0
            box_w = shape.width / EMU_PER_INCH * PX_PER_INCH
            box_h = shape.height / EMU_PER_INCH * PX_PER_INCH
            usable_w = box_w - (left + right) / EMU_PER_INCH * PX_PER_INCH
            usable_h = box_h - (top + bottom) / EMU_PER_INCH * PX_PER_INCH
            needed_h = 0.0
            for para in tf.paragraphs:
                text = para.text
                if not text:
                    continue
                para_w = estimate_paragraph_width(para)
                line_h = estimate_line_height(para)
                if usable_w > 1 and para_w > usable_w:
                    lines = ceil(para_w / usable_w)
                else:
                    lines = 1
                for ch in text:
                    if ch == "\n":
                        lines += 1
                needed_h += lines * line_h
                try:
                    if para.space_after is not None:
                        needed_h += para.space_after.pt * PX_PER_INCH / 72.0
                except Exception:
                    pass
                try:
                    if para.space_before is not None:
                        needed_h += para.space_before.pt * PX_PER_INCH / 72.0
                except Exception:
                    pass
            if usable_h > 1 and needed_h > usable_h * 1.04:
                flags.append((slide_idx, shape.shape_id, shape.name, box_w, box_h, usable_w, round(needed_h, 1), round(usable_h, 1), " | ".join(p.text for p in tf.paragraphs if p.text)[:70]))
    return flags


if __name__ == "__main__":
    for item in flag_shapes(sys.argv[1]):
        print("POSSIBLE", item)
    print("checked")
